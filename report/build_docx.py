"""Builds a Word version of the report from report/main.tex.

The LaTeX file stays the single source of truth: this script reads it,
renders every TikZ figure and displayed equation to PNG with pdfLaTeX, and
writes report/CaseStudyReport.docx (Times New Roman 12 pt, black and white,
A4, 1.5 line spacing). If report/sit_logo.png exists it is placed at the top
of the title page; otherwise a placeholder line is left for the logo.

Usage:  python report/build_docx.py      (needs pdflatex and pdftoppm)
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TEX = os.path.join(HERE, "main.tex")
OUT = os.path.join(HERE, "CaseStudyReport.docx")
LOGO = os.path.join(HERE, "sit_logo.png")
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

STANDALONE = r"""\documentclass[12pt,border=6pt]{standalone}
\usepackage[T1]{fontenc}
\usepackage{lmodern}
\usepackage{mathptmx}
\usepackage{amsmath}
\usepackage{xcolor}
\usepackage{tikz}
\usetikzlibrary{positioning,arrows.meta}
\usepackage{pgfplots}
\pgfplotsset{compat=1.17}
\newcommand{\attck}{ATT\&CK}
\begin{document}
%s
\end{document}
"""


# ----------------------------------------------------------------- helpers
def strip_comments(text: str) -> str:
    out = []
    for line in text.split("\n"):
        m = re.search(r"(?<!\\)%", line)
        if m:
            if line[: m.start()].strip() == "":
                continue
            line = line[: m.start()]
        out.append(line)
    return "\n".join(out)


def braced(text: str, i: int) -> tuple[str, int]:
    """Return the content of the {...} group starting at text[i] == '{'."""
    assert text[i] == "{", text[i:i + 30]
    depth = 0
    for j in range(i, len(text)):
        if text[j] == "{" and (j == 0 or text[j - 1] != "\\"):
            depth += 1
        elif text[j] == "}" and text[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return text[i + 1:j], j + 1
    raise ValueError("unbalanced braces")


def render_png(latex_body: str, workdir: str, name: str) -> str:
    tex = os.path.join(workdir, name + ".tex")
    with open(tex, "w") as fh:
        fh.write(STANDALONE % latex_body)
    subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", name + ".tex"],
                   cwd=workdir, check=True, stdout=subprocess.DEVNULL)
    subprocess.run(["pdftoppm", "-png", "-r", "300", "-singlefile", name + ".pdf", name],
                   cwd=workdir, check=True)
    return os.path.join(workdir, name + ".png")


def set_run_font(run, size=None, bold=None, italic=None, mono=False):
    run.font.name = "Courier New" if mono else FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Courier New" if mono else FONT)
    run.font.color.rgb = BLACK
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def fix_style_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in list(color.attrib):
            if attr.endswith("themeColor") or attr.endswith("themeShade") or attr.endswith("themeTint"):
                del color.attrib[attr]


# ------------------------------------------------------- inline conversion
class Ctx:
    def __init__(self):
        self.cite = {}
        self.ref = {}


MATH = {r"\times": "×", r"\le": "≤", r"\ge": "≥", r"\approx": "≈", r"\qquad": "  ",
        r"\quad": " ", r"\,": " ", r"\log": "log"}


def math_text(m: str) -> str:
    m = re.sub(r"\\text\{([^}]*)\}", r"\1", m)
    for k, v in MATH.items():
        m = m.replace(k, v)
    return m.replace("{", "").replace("}", "")


def compress(nums):
    nums = sorted(set(nums))
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else ", ".join(f"[{n}]" for n in nums[i:j + 1]))
        i = j + 1
    return ", ".join(parts)


def runs_of(text: str, ctx: Ctx, bold=False, italic=False, mono=False):
    """Convert inline LaTeX to a list of (text, bold, italic, mono) runs."""
    out = []
    buf = []

    def flush():
        if buf:
            out.append(("".join(buf), bold, italic, mono))
            buf.clear()

    i = 0
    while i < len(text):
        c = text[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+)\*?", text[i:])
            if not m:
                nxt = text[i + 1] if i + 1 < len(text) else ""
                if nxt == "\\":
                    buf.append("\n"); i += 2; continue
                if nxt in "%&_#$":
                    buf.append(nxt); i += 2; continue
                if nxt == ",":
                    buf.append("\u2009"); i += 2; continue
                if nxt == " ":
                    buf.append(" "); i += 2; continue
                i += 2
                continue
            name = m.group(1)
            j = i + len(m.group(0))
            arg = None
            if j < len(text) and text[j] == "{":
                arg, j = braced(text, j)
            if name in ("textbf", "head"):
                flush(); out += runs_of(arg, ctx, True, italic, mono)
                if name == "head":
                    out.append(("  ", bold, italic, mono))
            elif name in ("emph", "textit"):
                flush(); out += runs_of(arg, ctx, bold, True, mono)
            elif name == "texttt":
                flush(); out += runs_of(arg, ctx, bold, italic, True)
            elif name in ("textsc", "text", "mbox"):
                flush(); out += runs_of(arg, ctx, bold, italic, mono)
            elif name == "url":
                buf.append(arg)
            elif name == "cite":
                buf.append(compress([ctx.cite[k.strip()] for k in arg.split(",")]))
            elif name == "ref":
                buf.append(str(ctx.ref.get(arg, "?")))
            elif name == "attck":
                buf.append("ATT&CK")
            elif name in ("quad", "qquad"):
                buf.append(" ")
            else:  # \noindent, \centering, sizes, \par ... are layout only
                if arg is not None:
                    flush(); out += runs_of(arg, ctx, bold, italic, mono)
            i = j
            continue
        if c == "$":
            j = text.index("$", i + 1)
            flush()
            out.append((math_text(text[i + 1:j]), bold, True, mono))
            i = j + 1
            continue
        if c == "~":
            buf.append("\u00a0"); i += 1; continue
        if text.startswith("---", i):
            buf.append("—"); i += 3; continue
        if text.startswith("--", i):
            buf.append("–"); i += 2; continue
        if text.startswith("``", i):
            buf.append("“"); i += 2; continue
        if text.startswith("''", i):
            buf.append("”"); i += 2; continue
        if c == "`":
            buf.append("‘"); i += 1; continue
        if c in "{}":
            i += 1; continue
        if c == "\n":
            buf.append(" "); i += 1; continue
        buf.append(c)
        i += 1
    flush()
    # collapse whitespace inside runs
    cleaned = []
    for t, b, it, mo in out:
        t = re.sub(r"[ \t]+", " ", t)
        cleaned.append((t, b, it, mo))
    return cleaned


def add_runs(par, runs, size=12):
    for k, (t, b, it, mo) in enumerate(runs):
        if k == 0:
            t = t.lstrip()
        if not t:
            continue
        r = par.add_run(t)
        set_run_font(r, size=size, bold=b, italic=it, mono=mo)


# ------------------------------------------------------------ the builder
class Builder:
    def __init__(self, ctx: Ctx, workdir: str):
        self.ctx = ctx
        self.work = workdir
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.sec_n = 0
        self._setup()

    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.54))
        normal = d.styles["Normal"]
        fix_style_font(normal, 12)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(6)
        fix_style_font(d.styles["Heading 1"], 14, bold=True)
        h1 = d.styles["Heading 1"].paragraph_format
        h1.space_before, h1.space_after, h1.keep_with_next = Pt(18), Pt(8), True
        h1.line_spacing = 1.15

    # -- block helpers
    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, indent=None, spacing=None,
             after=None):
        p = self.doc.add_paragraph()
        p.alignment = align
        if spacing:
            p.paragraph_format.line_spacing = spacing
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if indent:
            p.paragraph_format.left_indent = Cm(indent[0])
            p.paragraph_format.first_line_indent = Cm(-indent[1])
        add_runs(p, runs_of(text, self.ctx), size)
        return p

    def heading(self, title, numbered):
        if numbered:
            self.sec_n += 1
            title = f"{self.sec_n}  {title}"
        h = self.doc.add_paragraph(style="Heading 1")
        r = h.add_run(title)
        set_run_font(r, size=14, bold=True)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def picture(self, png, max_in=6.2, scale=1.0):
        from PIL import Image
        w_px, _ = Image.open(png).size
        width = min(max_in, w_px / 300.0 * scale)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(png, width=Inches(width))
        return p

    def caption(self, kind, n, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(10 if kind == "Figure" else 4)
        if kind == "Table":
            p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{kind} {n}: ")
        set_run_font(r, bold=True)
        add_runs(p, runs_of(text, self.ctx))

    # -- environments
    def figure(self, body):
        cap = re.search(r"\\caption\{", body)
        caption, _ = braced(body, cap.end() - 1)
        tikz = re.search(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", body, re.S).group(0)
        tikz = re.sub(r"([\d.]+)\\textwidth", lambda m: f"{float(m.group(1)) * 15.92:.2f}cm", tikz)
        self.fig_n += 1
        png = render_png(tikz, self.work, f"fig{self.fig_n}")
        self.picture(png, scale=0.95)
        self.caption("Figure", self.fig_n, caption)

    def equation(self, body):
        self.eq_n += 1
        png = render_png("$\\displaystyle " + body.strip() + "$", self.work, f"eq{self.eq_n}")
        from PIL import Image
        w_px, _ = Image.open(png).size
        p = self.doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        p.add_run().add_picture(png, width=Inches(w_px / 300.0 * 0.95))
        r = p.add_run(f"\t({self.eq_n})")
        set_run_font(r)

    def table(self, body):
        cap = re.search(r"\\caption\{", body)
        caption, _ = braced(body, cap.end() - 1)
        m = re.search(r"\\begin\{(tabularx?)\}", body)
        j = m.end()
        if m.group(1) == "tabularx":
            _, j = braced(body, j)          # width argument
        spec, j = braced(body, j)
        inner = body[j:body.index("\\end{" + m.group(1) + "}")]
        cols = []
        widths_cm = {}
        k = 0
        while k < len(spec):
            ch = spec[k]
            if ch in "lcrX":
                cols.append(ch); k += 1
            elif ch in "pmb" and k + 1 < len(spec) and spec[k + 1] == "{":
                w, k = braced(spec, k + 1); cols.append("p"); widths_cm[len(cols) - 1] = float(w.rstrip("cm"))
            elif ch == ">" or ch == "<":
                _, k = braced(spec, k + 1)
            else:
                k += 1
        rows = []
        for raw in re.split(r"\\\\", inner):
            raw = re.sub(r"\\(toprule|midrule|bottomrule)(\[[^\]]*\])?", "", raw).strip()
            if not raw:
                continue
            rows.append([c.strip() for c in re.split(r"(?<!\\)&", raw)])
        total = 15.9
        for ci, c in enumerate(cols):
            if c in "lcr" and ci not in widths_cm:
                longest = max(len("".join(r for r, *_ in runs_of(row[ci], self.ctx))) if ci < len(row) else 0
                              for row in rows)
                widths_cm[ci] = min(6.5, 0.24 * longest + 0.6)
        flex = [ci for ci, c in enumerate(cols) if ci not in widths_cm]
        fixed = sum(widths_cm.values())
        for ci in flex:
            widths_cm[ci] = max(2.0, (total - fixed) / len(flex))
        scale = min(1.0, total / sum(widths_cm.values()))
        widths = [Cm(widths_cm[ci] * scale) for ci in range(len(cols))]
        self.tab_n += 1
        self.caption("Table", self.tab_n, caption)
        t = self.doc.add_table(rows=len(rows), cols=len(cols))
        t.autofit = False
        for ci, col in enumerate(t.columns):
            col.width = widths[ci]          # tblGrid/gridCol, read by LibreOffice
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            trpr = t.rows[ri]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for ci in range(len(cols)):
                cell = t.cell(ri, ci)
                cell.width = widths[ci]
                txt = row[ci] if ci < len(row) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = ri < len(rows) - 1   # keep table on one page
                p.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}.get(
                    cols[ci], WD_ALIGN_PARAGRAPH.LEFT)
                runs = runs_of(txt, self.ctx, bold=(ri == 0))
                add_runs(p, runs, size=12)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def enumerate(self, body):
        items = re.split(r"\\item\s*", body)[1:]
        for n, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f"{n}.\t")
            set_run_font(r)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            add_runs(p, runs_of(it.strip(), self.ctx))

    def bibliography(self, body):
        items = re.split(r"\\bibitem\{[^}]*\}\s*", body)[1:]
        for n, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            r = p.add_run(f"[{n}]\t")
            set_run_font(r)
            add_runs(p, runs_of(it.strip(), self.ctx))

    def title_page(self, body):
        """Each \\par-terminated chunk becomes one centred paragraph."""
        spacing = 0
        for chunk in re.split(r"\\par\b", body):
            chunk = chunk.strip()
            vs = re.findall(r"\\vspace\{([\d.]+)cm\}", chunk)
            for v in vs:
                spacing += float(v)
            if "\\vfill" in chunk:
                spacing += 2.5
            chunk = re.sub(r"\\vspace\{[^}]*\}|\\vfill|\\centering|\\setstretch\{[^}]*\}", "", chunk).strip()
            if not chunk:
                continue
            if "IfFileExists" in chunk:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if os.path.exists(LOGO):
                    p.add_run().add_picture(LOGO, height=Cm(3.2))
                else:
                    r = p.add_run("[SIT Pune logo – insert picture here]")
                    set_run_font(r, size=11, italic=True)
                continue
            size, bold = 12, None
            if "\\LARGE" in chunk:
                size, bold = 18, True
            elif "\\large" in chunk:
                size = 14
            chunk = re.sub(r"\\(LARGE|large|bfseries)\b", "", chunk).strip()
            while chunk.startswith("{") and chunk.count("{") > chunk.count("}"):
                chunk = chunk[1:].strip()
            while chunk.endswith("}") and chunk.count("}") > chunk.count("{"):
                chunk = chunk[:-1].strip()
            if not chunk:
                continue
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(min(spacing, 6.0) * 28.35 * 0.6)
            spacing = 0
            runs = runs_of(chunk, self.ctx, bold=bool(bold))
            add_runs(p, runs, size)

    # -- footer with page numbers on every page after the title page
    def finish_sections(self, first_section):
        sec = self.doc.sections[-1]
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi"):
            rf.set(qn(a), FONT)
        rpr.append(rf)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(pg)          # schema order: ... pgMar, pgNumType, cols ...
        else:
            sec._sectPr.append(pg)
        first_section.footer.is_linked_to_previous = False


# ------------------------------------------------------------------- main
def main():
    src = open(TEX).read()
    body = strip_comments(src.split("\\begin{document}")[1].split("\\end{document}")[0])
    ctx = Ctx()
    for n, key in enumerate(re.findall(r"\\bibitem\{([^}]*)\}", body), 1):
        ctx.cite[key] = n
    f = t = 0
    for env, label in re.findall(r"\\begin\{(figure|table)\}.*?\\label\{([^}]*)\}", body, re.S):
        if env == "figure":
            f += 1; ctx.ref[label] = f
        else:
            t += 1; ctx.ref[label] = t

    work = tempfile.mkdtemp()
    try:
        b = Builder(ctx, work)
        tp = re.search(r"\\begin\{titlepage\}(.*?)\\end\{titlepage\}", body, re.S)
        b.title_page(tp.group(1))
        first = b.doc.sections[0]
        b.doc.add_section(WD_SECTION.NEW_PAGE)
        b.finish_sections(first)
        rest = body[tp.end():]

        token = re.compile(
            r"\\section(\*?)\{([^}]*)\}"
            r"|\\begin\{(figure|table|equation|enumerate|thebibliography)\}(?:\[[^\]]*\])?(?:\{[^}]*\})?"
            r"|\\clearpage")
        pos = 0
        pending_break = False

        def emit_text(chunk):
            for par in re.split(r"\n\s*\n", chunk):
                par = re.sub(r"\\(setstretch|vspace)\{[^}]*\}", "", par).strip()
                par = par.replace("\\noindent", "").strip()
                if par:
                    b.para(par)

        while True:
            m = token.search(rest, pos)
            if not m:
                emit_text(rest[pos:]); break
            emit_text(rest[pos:m.start()])
            if m.group(0) == "\\clearpage":
                b.page_break(); pos = m.end(); continue
            if m.group(2) is not None:
                b.heading(m.group(2), numbered=(m.group(1) != "*"))
                pos = m.end(); continue
            env = m.group(3)
            end = rest.index("\\end{" + env + "}", m.end())
            inner = rest[m.end():end]
            if env == "figure":
                b.figure(inner)
            elif env == "table":
                b.table(inner)
            elif env == "equation":
                b.equation(inner)
            elif env == "enumerate":
                b.enumerate(inner)
            elif env == "thebibliography":
                b.heading("References", numbered=False)
                b.bibliography(inner)
            pos = end + len("\\end{" + env + "}")
        zoom = b.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        b.doc.core_properties.title = "AI-Powered Detection and Automated Response to Malicious PowerShell Attacks Using Machine Learning"
        b.doc.core_properties.author = "Kushagra"
        b.doc.save(OUT)
        print("wrote", OUT, "figures:", b.fig_n, "tables:", b.tab_n, "equations:", b.eq_n)
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
